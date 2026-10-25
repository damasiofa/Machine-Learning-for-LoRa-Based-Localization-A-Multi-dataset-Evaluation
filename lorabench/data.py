"""Dataset adapters, the common feature representation, spatial blocks, and the dataset audit.

Raw sources (not redistributed; see datasets/README.md):
  * Antwerp 2019 LoRaWAN dataset, Zenodo record 3904158 (v1.3)
  * Bonn LoRa dataset, GitHub mclab-hbrs/lora-bonn (commit 74d9e29)
  * Brno-Bucharest "LoRaWAN datasets 2021-2022", Zenodo record 7236698

Every scenario is converted once into a cache folder of .npy arrays (y, spatial_groups, rssi, snr,
esp, rx_mask, aggregates) plus metadata.json. All later stages read only the cache.
The parsing code is unchanged from the V5 benchmark (dataset_adapters_v5.py).
"""
from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path
from typing import Dict, Iterable, Iterator, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook

RSSI_MISSING = -200.0
SNR_MISSING = 0.0
ESP_MISSING = -200.0
SPATIAL_BLOCKS = 20
SPATIAL_BLOCK_RANDOM_STATE = 20260902

# Four source deployments with equal weight in the final synthesis.
DEPLOYMENTS = ["antwerp", "bonn", "brno", "bucharest"]
SCENARIOS_BY_DEPLOYMENT = {
    "antwerp": ["antwerp2019"],
    "bonn": ["bonn"],
    "brno": ["brno_iAG_BUT", "brno_iUG_BUT"],
    "bucharest": ["bucharest_iAG_UPB", "bucharest_o_UPB"],
}
PRIMARY_SCENARIOS = [s for d in DEPLOYMENTS for s in SCENARIOS_BY_DEPLOYMENT[d]]
ALL_DATASETS = list(PRIMARY_SCENARIOS)
COMMON_GROUPS = ["rssi", "snr", "esp", "rx_mask", "aggregates"]
COMMON_AUX_GROUPS = ["snr", "esp", "rx_mask", "aggregates"]

# Environment classes used in the article. Bucharest Outdoor is processed by the benchmark
# but excluded from the article's analysis (much smaller than the city-scale outdoor datasets
# and expressed in local Cartesian coordinates).
PAPER_OUTDOOR = ["antwerp2019", "bonn"]
PAPER_INDOOR = ["brno_iAG_BUT", "brno_iUG_BUT", "bucharest_iAG_UPB"]
PAPER_SCENARIOS = PAPER_OUTDOOR + PAPER_INDOOR
ENVIRONMENT_BY_SCENARIO = {**{s: "Outdoor" for s in PAPER_OUTDOOR}, **{s: "Indoor" for s in PAPER_INDOOR}}
SCENARIO_LABELS = {"antwerp2019": "Antwerp 2019", "bonn": "Bonn", "brno_iAG_BUT": "Brno iAG",
                   "brno_iUG_BUT": "Brno iUG", "bucharest_iAG_UPB": "Bucharest iAG",
                   "bucharest_o_UPB": "Bucharest Outdoor"}

DOWNLOAD_HINTS = {
    "antwerp": "Antwerp 2019: https://doi.org/10.5281/zenodo.3904158 (ZIP of the record, or the file "
               "lorawan_antwerp_2019_dataset.json.txt placed directly in the data folder)",
    "bonn": "Bonn: https://github.com/mclab-hbrs/lora-bonn (Code > Download ZIP; commit 74d9e29)",
    "brno": "Brno-Bucharest: https://zenodo.org/record/7236698 "
            "(LoRaWAN datasets 2021-2022-20221021T204201Z-001.zip)",
}


def deployment_for_scenario(dataset_id: str) -> str:
    for deployment, scenarios in SCENARIOS_BY_DEPLOYMENT.items():
        if dataset_id in scenarios:
            return deployment
    raise ValueError(dataset_id)


def expand_dataset_tokens(value: str) -> List[str]:
    v = value.strip()
    if v in ("primary", "all", "all_primary"):
        return list(PRIMARY_SCENARIOS)
    out: List[str] = []
    for token in [x.strip() for x in v.split(",") if x.strip()]:
        if token in SCENARIOS_BY_DEPLOYMENT:
            out.extend(SCENARIOS_BY_DEPLOYMENT[token])
        elif token in PRIMARY_SCENARIOS:
            out.append(token)
        else:
            raise ValueError(f"Unknown dataset/deployment: {token}; allowed={DEPLOYMENTS + PRIMARY_SCENARIOS}")
    # stable dedup
    return list(dict.fromkeys(out))


def _norm_name(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def discover_zip(data_dir: Path, kind: str) -> Path:
    zips = sorted(data_dir.glob("*.zip"))
    tokens = {
        "antwerp": ["3904158", "lorawan antwerp"],
        "bonn": ["lora bonn"],
        "brno": ["lorawan datasets 2021 2022"],
    }[kind]
    for z in zips:
        n = _norm_name(z.name)
        if any(t in n for t in tokens):
            return z
    raise FileNotFoundError(f"ZIP for {kind!r} not found in {data_dir}. Existing ZIPs: {[p.name for p in zips]}. "
                            f"Download from {DOWNLOAD_HINTS[kind]}")


def esp_from_rssi_snr(rssi: np.ndarray, snr: np.ndarray) -> np.ndarray:
    rssi = np.asarray(rssi, dtype=np.float64)
    snr = np.asarray(snr, dtype=np.float64)
    out = rssi - 10.0 * np.log10(1.0 + np.power(10.0, -snr / 10.0))
    return out.astype(np.float32)


def aggregate_radio_features(rssi: np.ndarray, snr: np.ndarray, esp: np.ndarray, mask: np.ndarray) -> np.ndarray:
    """Common 15-column packet-level radio summary, computed identically in every dataset."""
    n = rssi.shape[0]
    out = np.zeros((n, 15), dtype=np.float32)
    out[:, 0] = mask.sum(axis=1)
    col = 1
    for arr in (rssi, snr, esp):
        for i in range(n):
            vals = arr[i, mask[i] > 0.5]
            if len(vals):
                out[i, col:col+4] = [float(vals.mean()), float(vals.std()), float(vals.min()), float(vals.max())]
        col += 4
    for i in range(n):
        rv = rssi[i, mask[i] > 0.5]
        sv = snr[i, mask[i] > 0.5]
        out[i, 13] = float(rv.max() - rv.min()) if len(rv) else 0.0
        out[i, 14] = float(sv.max() - sv.min()) if len(sv) else 0.0
    return out


def _spatial_xy_m(y: np.ndarray, metric: str) -> np.ndarray:
    y = np.asarray(y, dtype=np.float64)
    if metric == "haversine_m":
        lat0 = np.radians(float(np.median(y[:, 0])))
        lat = np.radians(y[:, 0]); lon = np.radians(y[:, 1])
        lat_ref = np.radians(float(np.median(y[:, 0]))); lon_ref = np.radians(float(np.median(y[:, 1])))
        x = 6371000.0 * (lon - lon_ref) * np.cos(lat0)
        yy = 6371000.0 * (lat - lat_ref)
        return np.column_stack([x, yy])
    return y[:, :2].copy()


def _morton_codes_2d(xy: np.ndarray) -> np.ndarray:
    lo=xy.min(axis=0); hi=xy.max(axis=0); span=hi-lo; span[span<1e-12]=1.0
    q=np.clip(np.round((xy-lo)/span*65535.0),0,65535).astype(np.uint64)
    def spread(v):
        v=v & np.uint64(0x0000FFFF)
        v=(v | (v << np.uint64(8))) & np.uint64(0x00FF00FF)
        v=(v | (v << np.uint64(4))) & np.uint64(0x0F0F0F0F)
        v=(v | (v << np.uint64(2))) & np.uint64(0x33333333)
        v=(v | (v << np.uint64(1))) & np.uint64(0x55555555)
        return v
    return spread(q[:,0]) | (spread(q[:,1]) << np.uint64(1))


def common_spatial_groups(y: np.ndarray, metric: str, n_blocks: int = SPATIAL_BLOCKS) -> np.ndarray:
    """Build the same number of approximately equal-size spatial blocks in every scenario.

    Target positions are ordered by a 2-D Morton (Z-order) spatial code and cut into 20 contiguous,
    sample-balanced blocks. Repeated measurements at the exact same target position are never split.
    This gives every scenario the same 20-block holdout design without requiring geographic coordinates.
    """
    xy=_spatial_xy_m(y,metric)
    unique,inverse,counts=np.unique(xy,axis=0,return_inverse=True,return_counts=True)
    k=int(min(n_blocks,len(unique)))
    if k<3: raise ValueError(f"Need at least 3 unique spatial positions, found {len(unique)}")
    order=np.argsort(_morton_codes_2d(unique),kind="stable")
    sorted_counts=counts[order]; cum=np.cumsum(sorted_counts); total=int(cum[-1]); nu=len(unique)
    boundaries=[0]
    for j in range(1,k):
        target=j*total/k
        b=int(np.searchsorted(cum,target,side="left")+1)
        # keep at least one unique position for every remaining block
        b=max(b,boundaries[-1]+1); b=min(b,nu-(k-j))
        boundaries.append(b)
    boundaries.append(nu)
    unique_group=np.empty(nu,dtype=np.int32)
    for g,(a,b) in enumerate(zip(boundaries[:-1],boundaries[1:])):
        unique_group[order[a:b]]=g
    return unique_group[inverse].astype(np.int32)

def _save_cache(cache_dir: Path, arrays: Dict[str, np.ndarray], metadata: Dict) -> Path:
    cache_dir.mkdir(parents=True, exist_ok=True)
    for name, arr in arrays.items():
        np.save(cache_dir / f"{name}.npy", np.asarray(arr), allow_pickle=False)
    (cache_dir / "metadata.json").write_text(json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8")
    return cache_dir


def _stream_json_array(binary_file, chunk_size: int = 1 << 20) -> Iterator[dict]:
    """Streaming parser for a top-level JSON array without an external ijson dependency."""
    decoder = json.JSONDecoder()
    buf = ""; pos = 0; started = False; eof = False
    while True:
        if pos >= len(buf) - 1 and not eof:
            more = binary_file.read(chunk_size)
            if more:
                buf = buf[pos:] + more.decode("utf-8")
                pos = 0
            else:
                eof = True
        while pos < len(buf) and buf[pos].isspace(): pos += 1
        if not started:
            if pos >= len(buf):
                if eof: return
                continue
            if buf[pos] != "[": raise ValueError("Expected top-level JSON array")
            pos += 1; started = True
        while True:
            while pos < len(buf) and (buf[pos].isspace() or buf[pos] == ","): pos += 1
            if pos < len(buf) and buf[pos] == "]": return
            if pos >= len(buf): break
            try:
                obj, end = decoder.raw_decode(buf, pos)
            except json.JSONDecodeError:
                break
            yield obj
            pos = end
        if eof:
            tail = buf[pos:].strip()
            if tail in ("", "]"): return
            raise ValueError(f"Incomplete JSON near EOF: {tail[:100]!r}")
        more = binary_file.read(chunk_size)
        if not more:
            eof = True
        else:
            buf = buf[pos:] + more.decode("utf-8")
            pos = 0


class _PlainFile:
    """Minimal stand-in for ZipFile when the Antwerp JSON is stored unzipped."""
    def __init__(self, path: Path): self.path = Path(path)
    def open(self, _member): return self.path.open("rb")
    def __enter__(self): return self
    def __exit__(self, *exc): return False


def _antwerp_source(data_dir: Path):
    plain = sorted(Path(data_dir).glob("*lorawan_antwerp_2019_dataset.json.txt"))
    try:
        zpath = discover_zip(data_dir, "antwerp")
    except FileNotFoundError:
        if plain:
            return plain[0], _PlainFile(plain[0]), plain[0].name
        raise
    zf = zipfile.ZipFile(zpath)
    names = [n for n in zf.namelist() if n.lower().endswith("lorawan_antwerp_2019_dataset.json.txt")]
    if not names:
        zf.close()
        if plain:
            return plain[0], _PlainFile(plain[0]), plain[0].name
        raise FileNotFoundError("Antwerp JSON dataset not found")
    return zpath, zf, names[0]


def build_antwerp(data_dir: Path, cache_root: Path, force: bool = False) -> Path:
    cache = cache_root / "antwerp2019"
    if (cache / "metadata.json").exists() and not force:
        return cache
    zpath, zf, member = _antwerp_source(data_dir)
    with zf:
        gateways = set(); n = 0
        with zf.open(member) as f:
            for rec in _stream_json_array(f):
                if rec.get("latitude") is None or rec.get("longitude") is None: continue
                valid = False
                for g in (rec.get("gateways") or []):
                    if g.get("id") is None or g.get("rssi") is None or g.get("snr") is None: continue
                    gateways.add(str(g["id"])); valid = True
                if valid: n += 1
        gateway_ids = sorted(gateways); gi = {g:i for i,g in enumerate(gateway_ids)}; m = len(gateway_ids)
        rssi = np.full((n,m), RSSI_MISSING, dtype=np.float32)
        snr = np.full((n,m), SNR_MISSING, dtype=np.float32)
        mask = np.zeros((n,m), dtype=np.float32)
        y = np.empty((n,2), dtype=np.float64)
        sf_counts: Dict[str,int] = {}
        i = 0
        with zf.open(member) as f:
            for rec in _stream_json_array(f):
                if rec.get("latitude") is None or rec.get("longitude") is None: continue
                by: Dict[str,Tuple[float,float]] = {}
                for g in (rec.get("gateways") or []):
                    if g.get("id") is None or g.get("rssi") is None or g.get("snr") is None: continue
                    gid = str(g["id"]); rv = float(g["rssi"]); sv = float(g["snr"])
                    if gid not in by or rv > by[gid][0]: by[gid] = (rv,sv)
                if not by: continue
                y[i] = [float(rec["latitude"]), float(rec["longitude"])]
                for gid,(rv,sv) in by.items():
                    j=gi[gid]; rssi[i,j]=rv; snr[i,j]=sv; mask[i,j]=1.0
                sf = rec.get("sf"); sf_counts[str(sf)] = sf_counts.get(str(sf),0)+1
                i += 1
    esp=np.full_like(rssi,ESP_MISSING,dtype=np.float32); obs=mask>0.5; esp[obs]=esp_from_rssi_snr(rssi[obs],snr[obs])
    aggregates=aggregate_radio_features(rssi,snr,esp,mask)
    groups=common_spatial_groups(y,"haversine_m")
    arrays={"y":y,"spatial_groups":groups,"rssi":rssi,"snr":snr,"esp":esp,"rx_mask":mask,"aggregates":aggregates}
    metadata={
        "dataset_id":"antwerp2019","deployment_id":"antwerp","source_zip":zpath.name,"task_family":"localization_2d","target_metric":"haversine_m","target_coordinate_system":"WGS84 latitude/longitude",
        "n_samples":int(n),"n_gateways":int(m),"gateway_ids":gateway_ids,"sf_distribution":sf_counts,
        "feature_groups":{k:int(v.shape[1]) for k,v in arrays.items() if k not in ("y","spatial_groups")},
        "representations":{"RSSI_ONLY":["rssi"],"COMMON_FULL":list(COMMON_GROUPS)},
        "available_ablation_groups":list(COMMON_AUX_GROUPS),"spatial_holdout_supported":True,"spatial_block_method":f"Morton-order sample-balanced, {len(np.unique(groups))} blocks",
        "notes":["All valid SF 7-12 transmissions retained.","Only the 44 gateway IDs observed in at least one reception are represented.","ESP is recomputed from RSSI and SNR with the same formula used for every dataset; source-provided ESP is not used.","Packet metadata and gateway geometry are excluded from the equal-footing primary protocol."]
    }
    return _save_cache(cache,arrays,metadata)


def _extract_member(zf: zipfile.ZipFile, suffix: str, dest: Path) -> Path:
    matches=[n for n in zf.namelist() if n.replace("\\","/").endswith(suffix)]
    if not matches: raise FileNotFoundError(f"{suffix} not found inside archive")
    dest.parent.mkdir(parents=True,exist_ok=True)
    with zf.open(matches[0]) as src, dest.open("wb") as out: out.write(src.read())
    return dest


def build_bonn(data_dir: Path, cache_root: Path, force: bool = False) -> Path:
    cache=cache_root/"bonn"
    if (cache/"metadata.json").exists() and not force: return cache
    zpath=discover_zip(data_dir,"bonn"); extract=cache_root/"_sources"/"bonn"
    with zipfile.ZipFile(zpath) as zf:
        samples_p=_extract_member(zf,"survey/samples.csv",extract/"samples.csv")
    df=pd.read_csv(samples_p,dtype={"gw":str}).dropna(subset=["latitude","longitude","pkt_number","gw","rssi","snr"]).copy()
    keys=["transceived_at","latitude","longitude","pkt_number"]
    df=df.sort_values("rssi",ascending=False).drop_duplicates(keys+["gw"],keep="first")
    gateways=sorted(df["gw"].astype(str).unique())
    idx=pd.MultiIndex.from_frame(df[keys].drop_duplicates().sort_values(keys).reset_index(drop=True))
    def pivot(col): return df.pivot(index=keys,columns="gw",values=col).reindex(index=idx,columns=gateways)
    pr=pivot("rssi"); ps=pivot("snr")
    mask=pr.notna().to_numpy(dtype=np.float32); rssi=pr.fillna(RSSI_MISSING).to_numpy(dtype=np.float32); snr=ps.fillna(SNR_MISSING).to_numpy(dtype=np.float32)
    esp=np.full_like(rssi,ESP_MISSING,dtype=np.float32); obs=mask>0.5; esp[obs]=esp_from_rssi_snr(rssi[obs],snr[obs])
    y=np.column_stack([idx.get_level_values("latitude").astype(float),idx.get_level_values("longitude").astype(float)]).astype(np.float64)
    aggregates=aggregate_radio_features(rssi,snr,esp,mask); groups=common_spatial_groups(y,"haversine_m")
    arrays={"y":y,"spatial_groups":groups,"rssi":rssi,"snr":snr,"esp":esp,"rx_mask":mask,"aggregates":aggregates}
    metadata={
        "dataset_id":"bonn","deployment_id":"bonn","source_zip":zpath.name,"task_family":"localization_2d","target_metric":"haversine_m","target_coordinate_system":"WGS84 latitude/longitude",
        "n_samples":int(len(y)),"n_gateways":len(gateways),"gateway_ids":gateways,
        "feature_groups":{k:int(v.shape[1]) for k,v in arrays.items() if k not in ("y","spatial_groups")},
        "representations":{"RSSI_ONLY":["rssi"],"COMMON_FULL":list(COMMON_GROUPS)},"available_ablation_groups":list(COMMON_AUX_GROUPS),"spatial_holdout_supported":True,"spatial_block_method":f"Morton-order sample-balanced, {len(np.unique(groups))} blocks",
        "notes":["Transmission fingerprints reconstructed by date + GPS coordinate + packet number.","Gateway geometry is deliberately excluded from the equal-footing primary protocol."]
    }
    return _save_cache(cache,arrays,metadata)


def _json_from_excel_row(row: Sequence) -> Optional[Tuple[float,float,dict]]:
    start=next((i for i,v in enumerate(row) if isinstance(v,str) and "{" in v),None)
    if start is None: return None
    nums=[]
    for v in row[:start]:
        if v is None: continue
        try: nums.append(float(v))
        except Exception: pass
    if len(nums)>=3: x,y=nums[-2],nums[-1]
    elif len(nums)>=2: x,y=nums[0],nums[1]
    else: return None
    txt=",".join(str(v) for v in row[start:] if v is not None)
    try: obj=json.loads(txt)
    except Exception: return None
    return float(x),float(y),obj


def _extract_brno_raw(zpath: Path, cache_root: Path) -> Path:
    out=cache_root/"_sources"/"brno_raw"
    if out.exists() and list(out.rglob("*.xlsx")): return out
    out.mkdir(parents=True,exist_ok=True)
    with zipfile.ZipFile(zpath) as zf:
        for name in zf.namelist():
            n=name.replace("\\","/")
            if "/1. Raw option/" not in n or not n.lower().endswith(".xlsx"): continue
            if "/iUG_BUT/Top floor/" in n: continue
            rel=Path(n.split("LoRaWAN datasets 2021-2022/",1)[-1]); dest=out/rel; dest.parent.mkdir(parents=True,exist_ok=True)
            with zf.open(name) as src, dest.open("wb") as dst: dst.write(src.read())
    return out


def _scenario_files(raw_root: Path,dataset_id:str)->List[Path]:
    mapping={
        "brno_iAG_BUT":raw_root/"iAG_BUT/1. Raw option",
        "brno_iUG_BUT":raw_root/"iUG_BUT/Bottom floor/1. Raw option",
        "bucharest_iAG_UPB":raw_root/"iAG_UPB/1. Raw option",
        "bucharest_o_UPB":raw_root/"o_UPB/1. Raw option",
    }
    if dataset_id not in mapping: raise ValueError(dataset_id)
    return sorted(mapping[dataset_id].glob("*.xlsx"))


def build_brno_bucharest(dataset_id: str, data_dir: Path, cache_root: Path, force: bool=False) -> Path:
    cache=cache_root/dataset_id
    if (cache/"metadata.json").exists() and not force: return cache
    zpath=discover_zip(data_dir,"brno"); raw_root=_extract_brno_raw(zpath,cache_root); files=_scenario_files(raw_root,dataset_id)
    if not files: raise FileNotFoundError(f"No raw XLSX files for {dataset_id}")
    records=[]; gateways=set(); skipped=0
    for p in files:
        wb=load_workbook(p,read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]
        for row in ws.iter_rows(values_only=True):
            parsed=_json_from_excel_row(row)
            if parsed is None: continue
            x,y,obj=parsed; up=obj.get("uplink_message") or {}; rx=[]
            for r in (up.get("rx_metadata") or []):
                eui=(r.get("gateway_ids") or {}).get("eui")
                if not eui or r.get("rssi") is None or r.get("snr") is None: continue
                e=str(eui).upper(); rv=float(r["rssi"]); sv=float(r["snr"]); rx.append((e,rv,sv)); gateways.add(e)
            if rx: records.append((x,y,rx))
            else: skipped += 1
        wb.close()
    gids=sorted(gateways); gi={g:i for i,g in enumerate(gids)}; n=len(records); m=len(gids)
    rssi=np.full((n,m),RSSI_MISSING,dtype=np.float32); snr=np.full((n,m),SNR_MISSING,dtype=np.float32); mask=np.zeros((n,m),dtype=np.float32)
    y=np.empty((n,2),dtype=np.float64)
    for i,(xv,yv,rx) in enumerate(records):
        y[i]=[xv,yv]; by={}
        for e,rv,sv in rx:
            if e not in by or rv>by[e][0]: by[e]=(rv,sv)
        for e,(rv,sv) in by.items(): j=gi[e]; rssi[i,j]=rv; snr[i,j]=sv; mask[i,j]=1.0
    esp=np.full_like(rssi,ESP_MISSING,dtype=np.float32); obs=mask>0.5; esp[obs]=esp_from_rssi_snr(rssi[obs],snr[obs]); aggregates=aggregate_radio_features(rssi,snr,esp,mask)
    groups=common_spatial_groups(y,"euclidean_m")
    deployment=deployment_for_scenario(dataset_id); location="Brno, Czech Republic" if deployment=="brno" else "Bucharest, Romania"
    arrays={"y":y,"spatial_groups":groups,"rssi":rssi,"snr":snr,"esp":esp,"rx_mask":mask,"aggregates":aggregates}
    notes=["Only receptions with stable gateway EUI are used.",f"Rows with no identifiable EUI reception excluded: {skipped}.","Packet metadata is deliberately excluded from the equal-footing primary protocol."]
    if dataset_id=="brno_iUG_BUT": notes.append("Top-floor and Bottom-floor raw XLSX files are byte-identical; the Bottom-floor copy is read once. The target is effectively a 1-D corridor.")
    metadata={
        "dataset_id":dataset_id,"deployment_id":deployment,"source_zip":zpath.name,"task_family":"localization_2d","target_metric":"euclidean_m","target_coordinate_system":"local Cartesian x/y in meters","location":location,
        "n_samples":n,"n_gateways":m,"gateway_ids":gids,"feature_groups":{k:int(v.shape[1]) for k,v in arrays.items() if k not in ("y","spatial_groups")},
        "representations":{"RSSI_ONLY":["rssi"],"COMMON_FULL":list(COMMON_GROUPS)},"available_ablation_groups":list(COMMON_AUX_GROUPS),"spatial_holdout_supported":True,"spatial_block_method":f"Morton-order sample-balanced, {len(np.unique(groups))} blocks","notes":notes
    }
    return _save_cache(cache,arrays,metadata)


def build_dataset(dataset_id: str,data_dir:Path,cache_root:Path,force:bool=False)->Path:
    if dataset_id=="antwerp2019": return build_antwerp(data_dir,cache_root,force)
    if dataset_id=="bonn": return build_bonn(data_dir,cache_root,force)
    if dataset_id in {"brno_iAG_BUT","brno_iUG_BUT","bucharest_iAG_UPB","bucharest_o_UPB"}: return build_brno_bucharest(dataset_id,data_dir,cache_root,force)
    raise ValueError(dataset_id)


def build_many(dataset_ids: Sequence[str],data_dir:Path,cache_root:Path,force:bool=False)->Dict[str,Path]:
    return {d:build_dataset(d,data_dir,cache_root,force) for d in dataset_ids}


class FeatureStore:
    def __init__(self,cache_dir:Path):
        self.cache_dir=Path(cache_dir); self.meta=json.loads((self.cache_dir/"metadata.json").read_text(encoding="utf-8")); self.y=np.load(self.cache_dir/"y.npy",mmap_mode="r"); self.spatial_groups=np.load(self.cache_dir/"spatial_groups.npy",mmap_mode="r")
        self._groups={g:np.load(self.cache_dir/f"{g}.npy",mmap_mode="r") for g in self.meta["feature_groups"]}
    def groups_for_variant(self,variant:str)->List[str]:
        if variant in self.meta["representations"]: return list(self.meta["representations"][variant])
        if variant=="DROP_SNR_ESP": return [g for g in self.meta["representations"]["COMMON_FULL"] if g not in {"snr","esp"}]
        if variant.startswith("DROP_"):
            drop=variant[5:].lower()
            if drop not in COMMON_AUX_GROUPS: raise ValueError(variant)
            return [g for g in self.meta["representations"]["COMMON_FULL"] if g!=drop]
        raise ValueError(variant)
    def ablation_variants(self)->List[str]:
        return ["DROP_SNR","DROP_ESP","DROP_SNR_ESP","DROP_RX_MASK","DROP_AGGREGATES"]
    def get_x(self,variant:str,idx:np.ndarray)->np.ndarray:
        mats=[np.asarray(self._groups[g][idx]) for g in self.groups_for_variant(variant)]
        return np.concatenate(mats,axis=1).astype(np.float32,copy=False)
    def get_y(self,idx:np.ndarray)->np.ndarray: return np.asarray(self.y[idx],dtype=np.float64)


def audit(data_dir: Path, cache_root: Path, force: bool = False, datasets: Sequence[str] = None) -> bool:
    """Build (or reuse) the caches and print the audit of the common protocol (former auditar_datasets_v5.py)."""
    datasets = list(datasets or PRIMARY_SCENARIOS)
    data = Path(data_dir); cache = Path(cache_root)
    build_many(datasets, data, cache, force)
    print('=== V5 EQUAL-FOOTING DATASET AUDIT ===')
    print('Deployments (equal weight):', DEPLOYMENTS)
    print('Scenarios:', SCENARIOS_BY_DEPLOYMENT)
    common_ok = True
    for d in datasets:
        st = FeatureStore(cache/d); m = st.meta
        groups = list(m['representations']['COMMON_FULL']); common_ok &= (groups == COMMON_GROUPS)
        dims = {g: int(m['feature_groups'][g]) for g in COMMON_GROUPS}
        common_dim = sum(dims.values()); blocks = np.unique(st.spatial_groups); counts = np.bincount(np.asarray(st.spatial_groups, dtype=int))
        finite = True; variable = {}
        for g in COMMON_GROUPS:
            a = np.load(cache/d/f'{g}.npy', mmap_mode='r'); finite &= bool(np.isfinite(a).all()); variable[g] = int(np.sum(np.ptp(np.asarray(a, dtype=np.float64), axis=0) > 0))
        print(f'\n{d} -> deployment={m["deployment_id"]}')
        print(f'  n={m["n_samples"]} gateways={m["n_gateways"]} metric={m["target_metric"]}')
        print(f'  RSSI_ONLY={m["feature_groups"]["rssi"]} dims; COMMON_FULL={common_dim} dims')
        print(f'  common groups={groups}')
        print(f'  varying columns={variable}; all finite={finite}')
        print(f'  spatial blocks={len(blocks)}; block-size min/mean/max={counts.min()}/{counts.mean():.1f}/{counts.max()}')
        print(f'  ablations={st.ablation_variants()}')
        if len(blocks) != 20: common_ok = False
    print('\nCOMMON PROTOCOL COMPATIBLE:', common_ok)
    return bool(common_ok)
